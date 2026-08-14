"""00_Ingest.py — Convert raw source files to markdown with YAML frontmatter.

Walks data/files_raw/ for all supported file types, converts each to
markdown via docling (with pandas pre-processing for Excel), writes to
mirrored paths in data/files_ingested/.

YAML frontmatter is prepended to every output .md file with:
  - filename, source_file_path, last_converted, last_modified
  - filetype (original extension), language (langdetect)
  - entities (ticker lookup from entities.yaml)
  - sector (subsector from mapper.xlsx company→subsector lookup)
  - doctype (inferred from folder structure + filename patterns)

Outputs:
  - data/files_ingested/<mirrored>.md   — converted markdown files with frontmatter
  - data/files_ingested/manifest.json   — {md_rel_path: original_ext}
  - data/files_ingested/convert_hashes.json — SHA-256 for incremental

Run: python src/scripts/PMS2/00_Ingest.py
"""

from __future__ import annotations

import datetime
import hashlib
import json
import logging
import os
import re
import sys
import tempfile
import warnings
from contextlib import redirect_stderr
from io import BytesIO
from pathlib import Path

# ── Paths ────────────────────────────────────────────────────────────────
_PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent.parent

# Ensure project root on sys.path so `from src.xxx` imports work when
# this script is run directly (python src/scripts/PMS2/00_Ingest.py).
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

_RAW_DIR = _PROJECT_ROOT / "data" / "files_raw"
_INGESTED_DIR = _PROJECT_ROOT / "data" / "files_ingested"
_MANIFEST_PATH = _INGESTED_DIR / "manifest.json"
_HASHES_PATH = _INGESTED_DIR / "convert_hashes.json"

from src.scripts.PMS2.ingest_config import (  # noqa: E402
    SUPPORTED_EXTS, _EXCEL_EXTENSIONS, _DOCLING_FORMATS,
)

# ── Suppress noisy Docling/RapidOCR logging ──────────────────────────────
for _noisy in ("docling", "rapidocr", "RapidOCR", "tqdm"):
    logging.getLogger(_noisy).setLevel(logging.ERROR)
logging.getLogger("docling.pipeline").setLevel(logging.ERROR)
os.environ.setdefault("RAPIDOCR_LOG_LEVEL", "ERROR")

warnings.filterwarnings("ignore", message=".*text detection result is empty.*")
warnings.filterwarnings("ignore", message=".*RapidOCR returned empty result.*")


# ═════════════════════════════════════════════════════════════════════════
# Hash helpers
# ═════════════════════════════════════════════════════════════════════════

def _sha256_file(path: Path) -> str | None:
    """Compute SHA-256 hex digest. Returns None if file is missing/unreadable."""
    try:
        h = hashlib.sha256()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                h.update(chunk)
        return h.hexdigest()
    except (FileNotFoundError, PermissionError, OSError):
        return None


def _load_hashes() -> dict[str, str]:
    try:
        return json.loads(_HASHES_PATH.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return {}


# ═════════════════════════════════════════════════════════════════════════
# File discovery
# ═════════════════════════════════════════════════════════════════════════

def _find_raw_files() -> list[tuple[Path, str]]:
    """Find all supported files in files_raw/.
    Returns list of (absolute_path, relative_path_str) tuples.
    """
    files = []
    for path in sorted(_RAW_DIR.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTS:
            rel = str(path.relative_to(_RAW_DIR))
            files.append((path, rel))
    return files


def _md_rel_path(raw_rel: str) -> str:
    """Convert raw relative path to ingested .md relative path."""
    p = Path(raw_rel)
    return str(p.with_suffix(".md"))


# ═════════════════════════════════════════════════════════════════════════
# Mapper — company → subsector lookup from mapper.xlsx
# ═════════════════════════════════════════════════════════════════════════

_mapper_cache: dict[str, list[str]] | None = None


def _load_mapper(mapper_path: str = "mapper.xlsx") -> dict[str, list[str]]:
    """Parse mapper.xlsx and return {company_lower: [subsector, ...]}.
    Cached in memory — only loaded once per session.
    """
    global _mapper_cache
    if _mapper_cache is not None:
        return _mapper_cache

    mp = _PROJECT_ROOT / mapper_path
    if not mp.exists():
        _mapper_cache = {}
        return _mapper_cache

    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(mp), data_only=True)
        ws = wb.active

        subsectors: list[str] = []
        for col in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=col).value
            subsectors.append(str(val).strip().lower() if val else "")

        company_map: dict[str, list[str]] = {}
        for row in range(2, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                val = ws.cell(row=row, column=col).value
                if val and str(val).strip():
                    company = str(val).strip()
                    sub = subsectors[col - 1] if col - 1 < len(subsectors) else ""
                    if sub:
                        company_lower = company.lower()
                        company_map.setdefault(company_lower, []).append(sub)

        wb.close()
        _mapper_cache = company_map
        return company_map
    except Exception:
        _mapper_cache = {}
        return _mapper_cache


def _infer_company_from_path(rel_path: str) -> str | None:
    """Infer the company name from a file's relative path within data/.

    Examples:
        ADI/ADI FY18 10-K.pdf          → ADI
        LITE/Notes/lite.txt            → LITE
        0 IPO/JD Health/some.pdf       → JD Health
        Meituan/Model/ms model.xlsm    → Meituan
    """
    parts = Path(rel_path).parts
    if not parts:
        return None

    top = parts[0]
    if top == "0 IPO" and len(parts) >= 3:
        return parts[1]
    elif top == "0 IPO" and len(parts) == 2:
        return "0 IPO"
    else:
        return top


def _resolve_tags(company: str | None, mapper: dict[str, list[str]]) -> list[str]:
    """Given a company name, return [company] + [subsector, ...] tags."""
    tags: list[str] = []
    if company:
        tags.append(company)
        company_lower = company.lower()
        if company_lower in mapper:
            tags.extend(mapper[company_lower])
        else:
            for mapper_key, subsectors in mapper.items():
                if company_lower in mapper_key or mapper_key in company_lower:
                    tags.extend(subsectors)
                    break
    return tags


# ═════════════════════════════════════════════════════════════════════════
# Docling converter (cached singleton)
# ═════════════════════════════════════════════════════════════════════════

_docling_converter: "DocumentConverter | None" = None


def _get_docling_converter() -> "DocumentConverter":
    global _docling_converter
    if _docling_converter is None:
        from docling.document_converter import DocumentConverter
        _docling_converter = DocumentConverter()
    return _docling_converter


# ═════════════════════════════════════════════════════════════════════════
# Pandas Excel preprocessing
# ═════════════════════════════════════════════════════════════════════════

def _preprocess_excel_with_pandas(source: Path) -> tuple[Path, Path | None]:
    """Drop entirely empty columns from an Excel file before Docling conversion.

    Returns ``(source_to_use, temp_file_to_cleanup)``.  When no columns are
    removed the original ``source`` is returned with a ``None`` cleanup handle.
    Gracefully degrades if pandas/openpyxl are unavailable.
    """
    try:
        import pandas as pd
    except ImportError:
        return (source, None)

    try:
        xl = pd.read_excel(source, sheet_name=None, header=None)
    except Exception:
        return (source, None)

    cleaned_sheets: dict = {}
    has_changes = False

    for sheet_name, df in xl.items():
        original_cols = len(df.columns)
        df = df.dropna(axis=1, how="all")

        cols_to_drop = []
        for col in df.columns:
            non_empty = df[col].apply(
                lambda x: not (
                    pd.isna(x)
                    or (isinstance(x, str) and x.strip() == "")
                )
            ).sum()
            if non_empty == 0:
                cols_to_drop.append(col)

        if cols_to_drop:
            df = df.drop(columns=cols_to_drop)

        if len(df.columns) < original_cols or cols_to_drop:
            has_changes = True

        if not df.empty:
            cleaned_sheets[sheet_name] = df

    if not has_changes:
        return (source, None)

    if not cleaned_sheets:
        return (source, None)

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp_path = Path(tmp.name)
    tmp.close()

    try:
        with pd.ExcelWriter(str(tmp_path), engine="openpyxl") as writer:
            for sheet_name, df in cleaned_sheets.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False, header=False)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        return (source, None)

    return (tmp_path, tmp_path)


# ═════════════════════════════════════════════════════════════════════════
# Conversion
# ═════════════════════════════════════════════════════════════════════════

def _convert_with_docling(source: Path) -> tuple[str, int]:
    """Convert a file to Markdown using Docling.

    Handles PDF, DOCX, PPTX, XLSX, XLSM, HTML, CSV, MD, TXT, and images.
    Returns (markdown_text, page_count).

    Falls back to DocumentStream when Docling's internal filetype detection
    fails on long/Unicode Windows paths.
    """
    from docling.datamodel.base_models import DocumentStream
    from docling.document_converter import ConversionError

    dc = _get_docling_converter()

    with open(os.devnull, "w") as devnull, redirect_stderr(devnull):
        try:
            result = dc.convert(source)
        except ConversionError:
            try:
                data = source.read_bytes()
                stream = DocumentStream(
                    name=source.name,
                    stream=BytesIO(data),
                )
                result = dc.convert(stream)
            except Exception:
                raise

    md = result.document.export_to_markdown()

    try:
        page_count = result.document.num_pages()
    except Exception:
        page_count = 1

    md = _clean_picture_placeholders(md)
    md = _clean_markdown_tables(md)

    return md, page_count


def _convert_with_markitdown(source: Path) -> tuple[str, int]:
    """Convert a file to Markdown using MarkItDown (fallback for .xml, .json)."""
    from markitdown import MarkItDown
    md = MarkItDown()
    result = md.convert(str(source))
    return result.text_content, 1


# ═════════════════════════════════════════════════════════════════════════
# Markdown post-processing
# ═════════════════════════════════════════════════════════════════════════

def _clean_picture_placeholders(md_text: str) -> str:
    """Strip auto-generated picture placeholder markers from Docling output."""
    md_text = re.sub(
        r'\*{0,2}==> picture \[\d+ x \d+\] intentionally omitted <==\*{0,2}',
        '', md_text,
    )
    md_text = re.sub(
        r'\*{0,2}-{5} Start of picture text -{5}\*{0,2}<br>',
        '', md_text,
    )
    md_text = re.sub(
        r'<br>\*{0,2}-{5} End of picture text -{5}\*{0,2}',
        '', md_text,
    )
    md_text = re.sub(r'\n\s*-{5,}\s*\n', '\n\n', md_text)
    md_text = re.sub(r'\n{3,}', '\n\n', md_text)
    return md_text


def _clean_markdown_tables(md_text: str) -> str:
    """Post-process markdown to fix common PDF table artifacts.

    Fixes applied to each pipe-table block:
      1. Remove entirely empty columns.
      2. Collapse duplicate columns (identical text across all rows).
      3. Regenerate separator row to match final column count.
    """
    lines = md_text.split('\n')
    result: list[str] = []
    i = 0

    while i < len(lines):
        line = lines[i]

        if not line.strip().startswith('|'):
            result.append(line)
            i += 1
            continue

        table_lines: list[str] = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            table_lines.append(lines[i])
            i += 1

        if len(table_lines) < 2:
            result.extend(table_lines)
            continue

        cleaned = _clean_single_table(table_lines)
        result.extend(cleaned)

    return '\n'.join(result)


def _clean_single_table(table_lines: list[str]) -> list[str]:
    """Clean a single pipe-table block. Returns cleaned lines."""
    parsed: list[list[str]] = []
    for line in table_lines:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        parsed.append(cells)

    if not parsed:
        return table_lines

    max_cols = max(len(row) for row in parsed)
    for row in parsed:
        while len(row) < max_cols:
            row.append('')

    # 1. Identify entirely empty columns
    keep_cols: list[int] = []
    for col_idx in range(max_cols):
        has_content = any(row[col_idx].strip() for row in parsed)
        if has_content:
            keep_cols.append(col_idx)

    if not keep_cols:
        return []

    # 2. Collapse duplicate columns (don't collapse first column = row labels)
    final_cols: list[int] = [keep_cols[0]]
    for col_idx in keep_cols[1:]:
        is_duplicate = False
        for existing_idx in final_cols:
            if all(
                row[col_idx].strip() == row[existing_idx].strip()
                for row in parsed
            ):
                is_duplicate = True
                break
        if not is_duplicate:
            final_cols.append(col_idx)

    if final_cols == list(range(max_cols)):
        return table_lines

    # 3. Rebuild table with pruned columns
    result: list[str] = []
    has_separator = any(
        all('---' in c or not c.strip() for c in row)
        for row in parsed
    )

    for row in parsed:
        is_sep = all('---' in c or not c.strip() for c in row)
        if is_sep:
            result.append('| ' + ' | '.join(['---'] * len(final_cols)) + ' |')
        else:
            cells = [row[c] for c in final_cols]
            result.append('| ' + ' | '.join(cells) + ' |')

    if not has_separator and len(result) > 1:
        sep = '| ' + ' | '.join(['---'] * len(final_cols)) + ' |'
        result.insert(1, sep)

    return result


# ═════════════════════════════════════════════════════════════════════════
# .msg email conversion
# ═════════════════════════════════════════════════════════════════════════

def _convert_msg_file(source: Path) -> tuple[str, int]:
    """Convert an Outlook .msg file to markdown.

    Extracts metadata, converts HTML body via markdownify (fallback to
    plain text), and returns (markdown_text, page_count).
    """
    import extract_msg
    import markdownify as mdify

    try:
        msg = extract_msg.Message(str(source))
    except Exception:
        return ("", 0)

    try:
        # ── Metadata ───────────────────────────────────────────────────
        sender = str(getattr(msg, "sender", "") or "")
        to = str(getattr(msg, "to", "") or "")
        cc = str(getattr(msg, "cc", "") or "")
        subject = str(getattr(msg, "subject", "") or "(No Subject)")
        date_val = getattr(msg, "date", None)

        date_str = ""
        if date_val is not None:
            try:
                date_str = date_val.strftime("%Y-%m-%d %H:%M UTC")
            except Exception:
                date_str = str(date_val)

        # ── Body: try HTML first, fall back to plain text ───────────────
        body_md = ""
        html_body = getattr(msg, "htmlBody", None)
        if html_body:
            if isinstance(html_body, bytes):
                html_body = html_body.decode("utf-8", errors="replace")
            if html_body and html_body.strip():
                try:
                    # Strip Outlook/Word CSS comments and Office tags
                    html_body = re.sub(r"<!--.*?-->", "", html_body, flags=re.DOTALL)
                    html_body = re.sub(r"</?(?:o:|w:|v:|m:|st1:)[^>]*>", "", html_body, flags=re.IGNORECASE)
                    html_body = re.sub(r"<style[^>]*>.*?</style>", "", html_body, flags=re.DOTALL | re.IGNORECASE)
                    body_md = mdify.markdownify(
                        html_body, heading_style="ATX", bullets="-",
                        strip=["script", "style"],
                    )
                except Exception:
                    body_md = ""

        if not body_md or not body_md.strip():
            plain_body = getattr(msg, "body", None)
            if plain_body and str(plain_body).strip():
                body_md = str(plain_body)

        if not body_md or not body_md.strip():
            body_md = "*(No body content)*"

        # ── Build email header ─────────────────────────────────────────
        header_lines = [f"# {subject}", ""]
        header_lines.append("| | |")
        header_lines.append("|---|---|")
        if sender:
            header_lines.append(f"| **From** | {sender} |")
        if to:
            header_lines.append(f"| **To** | {to} |")
        if date_str:
            header_lines.append(f"| **Date** | {date_str} |")
        if cc:
            header_lines.append(f"| **Cc** | {cc} |")
        header_lines.append("")
        header_lines.append("---")
        header_lines.append("")

        md_text = "\n".join(header_lines) + body_md.rstrip() + "\n"

        return (md_text, 1)

    finally:
        try:
            msg.close()
        except Exception:
            pass


# ── Path shortening for Windows MAX_PATH ─────────────────────────────────
_MAX_PATH = 250  # stay under Windows 260-char limit with margin


def _shorten_output_path(md_abs: Path) -> Path:
    """Shorten an output path if it exceeds Windows MAX_PATH.

    Truncates the filename stem (keeping the .md extension) to fit within
    the limit.  Returns the original path if it's already short enough.
    """
    path_str = str(md_abs)
    if len(path_str) <= _MAX_PATH:
        return md_abs

    # Shorten filename stem, preserving extension and parent directory
    parent = md_abs.parent
    stem = md_abs.stem
    suffix = md_abs.suffix  # .md

    max_stem = _MAX_PATH - len(str(parent)) - len(suffix) - 1  # -1 for separator
    if max_stem < 20:
        max_stem = 20  # absolute minimum to keep it identifiable

    if len(stem) > max_stem:
        stem = stem[:max_stem].rstrip("_").rstrip(". ").rstrip("-")

    shortened = parent / f"{stem}{suffix}"
    return shortened


# ═════════════════════════════════════════════════════════════════════════
# YAML frontmatter
# ═════════════════════════════════════════════════════════════════════════

def _build_frontmatter(
    *,
    source_path: str,
    source_type: str,
    tags: list[str] | None = None,
    input_dir: Path | None = None,
    entities_map: dict[str, list[str]] | None = None,
    keywords_config: dict | None = None,
    md_text: str = "",
) -> str:
    """Build YAML frontmatter with metadata fields."""
    from src.scripts.PMS2.frontmatter_helpers import (
        detect_doctype, detect_entities, detect_language,
        get_file_mtime, yaml_list,
    )

    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    company = _infer_company_from_path(source_path)

    file_ext = Path(source_path).suffix if source_path else ""
    filename = Path(source_path).name if source_path else ""
    abs_path = str((_PROJECT_ROOT / "data" / "files_raw" / source_path).resolve()) if source_path else ""

    last_modified = ""
    if input_dir:
        last_modified = get_file_mtime(input_dir, source_path)

    language = detect_language(md_text)
    entities = detect_entities(company, entities_map or {})
    sector = tags[1:] if tags and len(tags) > 1 else []
    doctype = detect_doctype(source_path, file_ext, keywords_config)
    doctype = [] if doctype == ["unknown"] else doctype

    lines = [
        "---",
        f"filename: {filename}",
        f"source_file_path: {abs_path}",
        f"last_converted: {timestamp}",
        f"last_modified: {last_modified}",
        f"filetype: {source_type}",
        f"language: {yaml_list(language)}",
        f"entities: {yaml_list(entities)}",
        f"sector: {yaml_list(sector)}",
        f"doctype: {yaml_list(doctype)}",
        "---",
        "",
    ]
    return "\n".join(lines)


# ═════════════════════════════════════════════════════════════════════════
# Orphan cleanup
# ═════════════════════════════════════════════════════════════════════════

def _cleanup_orphans(input_dir: Path, output_dir: Path) -> int:
    """Remove .md files in output_dir whose source file no longer exists.

    Returns the number of orphaned files removed.
    """
    import frontmatter as fm

    if not output_dir.is_dir():
        return 0

    removed = 0
    empty_dirs: set[Path] = set()

    for md_file in sorted(output_dir.rglob("*.md")):
        source_abs = None
        try:
            post = fm.load(str(md_file))
            sp = post.get("source_file_path")
            if sp and isinstance(sp, str):
                source_abs = sp
        except Exception:
            pass

        if not source_abs:
            rel_md = md_file.relative_to(output_dir)
            for suffix in SUPPORTED_EXTS:
                candidate = input_dir / str(rel_md).replace(".md", suffix)
                if candidate.exists():
                    source_abs = str(candidate)
                    break

        if source_abs and Path(source_abs).exists():
            continue

        try:
            md_file.unlink()
            removed += 1
            print(f"  orphan: {md_file.relative_to(output_dir)}")
            p = md_file.parent
            while p != output_dir:
                empty_dirs.add(p)
                p = p.parent
        except Exception as exc:
            print(f"  could not remove {md_file.relative_to(output_dir)}: {exc}")

    for d in sorted(empty_dirs, key=lambda x: len(str(x)), reverse=True):
        try:
            if d.is_dir() and not any(d.iterdir()):
                d.rmdir()
        except OSError:
            pass

    if removed:
        print(f"\nRemoved {removed} orphaned .md file(s) (source deleted).\n")

    return removed


# ═════════════════════════════════════════════════════════════════════════
# Main pipeline
# ═════════════════════════════════════════════════════════════════════════

def run_ingest() -> None:
    """Main entry point for 00_Ingest."""
    # 1. Find all supported raw files
    raw_files = _find_raw_files()
    if not raw_files:
        print("[00_Ingest] No supported files found in files_raw/.")
        return

    print(f"[00_Ingest] Found {len(raw_files)} supported files in {_RAW_DIR}")

    # 2. Load existing hashes for incremental
    old_hashes = _load_hashes()
    new_hashes: dict[str, str] = {}

    _INGESTED_DIR.mkdir(parents=True, exist_ok=True)

    # 3. Load mapper, entities, keywords config
    mapper = _load_mapper()
    if mapper:
        print(f"  Mapper: {len(mapper)} companies loaded.")

    from src.scripts.PMS2.frontmatter_helpers import load_entities, load_keywords_config
    entities_map = load_entities()
    keywords_config = load_keywords_config()

    # 4. Docling's httpx picks up ALL_PROXY (SOCKS) — suppress
    _saved_all_proxy = os.environ.pop("ALL_PROXY", None)
    _saved_all_proxy_lc = os.environ.pop("all_proxy", None)

    converted = 0
    skipped = 0
    failed = 0

    for abs_path, raw_rel in raw_files:
        file_hash = _sha256_file(abs_path)
        if file_hash is None:
            print(f"  Skipped (missing/inaccessible): {raw_rel}")
            skipped += 1
            continue
        new_hashes[raw_rel] = file_hash

        md_rel = _md_rel_path(raw_rel)
        md_abs = _INGESTED_DIR / md_rel

        # Skip if hash unchanged AND output .md exists
        if raw_rel in old_hashes and old_hashes[raw_rel] == file_hash and md_abs.exists():
            skipped += 1
            continue

        suffix = abs_path.suffix.lower()

        try:
            # ── Convert ────────────────────────────────────────────────
            if suffix in _DOCLING_FORMATS:
                conv_source = abs_path
                tmp_excel = None
                if suffix in _EXCEL_EXTENSIONS:
                    conv_source, tmp_excel = _preprocess_excel_with_pandas(abs_path)
                try:
                    md_text, page_count = _convert_with_docling(conv_source)
                finally:
                    if tmp_excel:
                        tmp_excel.unlink(missing_ok=True)
            elif suffix == ".msg":
                md_text, page_count = _convert_msg_file(abs_path)
                if not md_text or len(md_text.strip()) < 50:
                    print(f"  Skipped (empty .msg): {raw_rel}")
                    skipped += 1
                    continue
            elif suffix in (".xml", ".json"):
                md_text, page_count = _convert_with_markitdown(abs_path)
            else:
                md_text, page_count = _convert_with_markitdown(abs_path)

            # Skip near-empty output
            if len(md_text.strip()) < 50:
                print(f"  Skipped (near-empty): {raw_rel}")
                skipped += 1
                continue

            # ── Build frontmatter ──────────────────────────────────────
            company = _infer_company_from_path(raw_rel)
            tags = _resolve_tags(company, mapper)

            frontmatter = _build_frontmatter(
                source_path=raw_rel,
                source_type=suffix.lstrip("."),
                tags=tags,
                input_dir=_RAW_DIR,
                entities_map=entities_map,
                keywords_config=keywords_config,
                md_text=md_text,
            )

            # ── Write output ───────────────────────────────────────────
            md_abs = _shorten_output_path(md_abs)
            md_abs.parent.mkdir(parents=True, exist_ok=True)
            md_abs.write_text(frontmatter + md_text, encoding="utf-8")
            converted += 1
            tag_display = f"  tags: {tags}" if tags else ""
            print(f"  Converted: {raw_rel}{tag_display}")

        except PermissionError:
            print(f"  Skipped (OneDrive cloud-only): {raw_rel}")
            skipped += 1
        except Exception as exc:
            msg = f"{type(exc).__name__}: {exc}"
            failed += 1
            print(f"  Failed: {raw_rel} — {msg}")

    # Restore proxy env vars
    if _saved_all_proxy is not None:
        os.environ["ALL_PROXY"] = _saved_all_proxy
    if _saved_all_proxy_lc is not None:
        os.environ["all_proxy"] = _saved_all_proxy_lc

    print(f"\n[00_Ingest] {converted} converted, {skipped} unchanged, {failed} failed")

    # 5. Build manifest from ALL raw files (not just converted)
    manifest: dict[str, str] = {}
    for _, raw_rel in raw_files:
        md_rel = _md_rel_path(raw_rel)
        ext = Path(raw_rel).suffix.lstrip(".").lower()
        manifest[md_rel] = ext

    _MANIFEST_PATH.write_text(
        json.dumps(manifest, indent=2, sort_keys=True, ensure_ascii=False),
        encoding="utf-8",
    )
    print(f"[00_Ingest] manifest.json: {len(manifest)} entries")

    # 6. Write hashes
    _HASHES_PATH.write_text(
        json.dumps(new_hashes, indent=2, sort_keys=True, ensure_ascii=False),
        encoding="utf-8",
    )
    print(f"[00_Ingest] convert_hashes.json written")

    # 7. Clean up orphans
    _cleanup_orphans(_RAW_DIR, _INGESTED_DIR)


if __name__ == "__main__":
    run_ingest()
